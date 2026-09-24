#!/usr/bin/env python3
"""
Shared plumbing for the Alien Tiles SAT variants.

Instance loading, CLI parsing, printing, verification, timeouts, Excel export.
The encoding and the solving strategies live in encoder.py and solver.py.
"""

from __future__ import annotations

import argparse
import glob
import json
import os
import platform
import subprocess
import sys
import time
from contextlib import contextmanager
from dataclasses import dataclass, field
from multiprocessing import Pipe, Process

import openpyxl
from openpyxl.styles import Font


# ══════════════════════════════════════════════════════════════════════════
#  Result container
# ══════════════════════════════════════════════════════════════════════════

@dataclass
class SolveResult:
    """What a variant hands back to the plumbing."""
    X: list[list[int]] | None = None           # click matrix, None if UNSAT
    optimum: int | None = None                 # total clicks, if minimising
    n_vars: int = 0
    n_clauses: int = 0
    extra: dict = field(default_factory=dict)


# ══════════════════════════════════════════════════════════════════════════
#  Instance I/O
# ══════════════════════════════════════════════════════════════════════════

def load_instance(path: str) -> dict:
    """Load one instance JSON. Name falls back to the filename stem."""
    with open(path) as f:
        data = json.load(f)
    name = data.get("name") or os.path.splitext(os.path.basename(path))[0]
    return {"name": name, "N": data["N"], "c": data["c"], "target": data["target"]}


def find_instance_files(dirpath: str) -> list[str]:
    """Collect instance JSONs under `dirpath` — anything carrying N, c and target."""
    if not os.path.isdir(dirpath):
        print(f"Not a directory: {dirpath}")
        sys.exit(1)

    files = sorted(glob.glob(os.path.join(dirpath, "**", "*.json"), recursive=True))
    instances = []
    for filepath in files:
        try:
            with open(filepath) as f:
                data = json.load(f)
        except (json.JSONDecodeError, OSError):
            continue
        if isinstance(data, dict) and all(k in data for k in ("N", "c", "target")):
            instances.append(filepath)

    if not instances:
        print(f"No instance .json files (with N/c/target) found in {dirpath}")
        sys.exit(1)
    return instances


def parse_target(s: str, N: int, c: int) -> list[list[int]]:
    """Parse a CLI target string "v,v,...;v,v,..." into an N×N int matrix."""
    rows = s.strip().split(";")
    if len(rows) != N:
        raise ValueError(f"Expected {N} rows, got {len(rows)}")

    matrix = []
    for ri, row_str in enumerate(rows):
        vals = [int(x.strip()) for x in row_str.split(",")]
        if len(vals) != N:
            raise ValueError(f"Row {ri}: expected {N} values, got {len(vals)}")
        for v in vals:
            if not 0 <= v < c:
                raise ValueError(f"Row {ri}: value {v} outside [0, {c - 1}]")
        matrix.append(vals)
    return matrix


# ══════════════════════════════════════════════════════════════════════════
#  Provenance
# ══════════════════════════════════════════════════════════════════════════

_MACHINE = None


def machine_label() -> str:
    """CPU and OS of the machine producing a row, e.g. "Apple M1 | Darwin 25.6.0"."""
    global _MACHINE
    if _MACHINE is not None:
        return _MACHINE

    cpu = ""
    try:
        if platform.system() == "Darwin":
            cpu = subprocess.run(["sysctl", "-n", "machdep.cpu.brand_string"],
                                 capture_output=True, text=True,
                                 timeout=5).stdout.strip()
        elif platform.system() == "Linux":
            with open("/proc/cpuinfo") as f:
                for line in f:
                    if line.startswith("model name"):
                        cpu = line.split(":", 1)[1].strip()
                        break
    except (OSError, subprocess.SubprocessError):
        cpu = ""
    cpu = cpu or platform.processor() or platform.machine() or "unknown CPU"

    _MACHINE = f"{cpu} | {platform.system()} {platform.release()}"
    return _MACHINE


# ══════════════════════════════════════════════════════════════════════════
#  Printing and verification
# ══════════════════════════════════════════════════════════════════════════

def print_board(board: list[list[int]], title: str = "Board"):
    print(f"{title}:")
    for row in board:
        print("  " + " ".join(str(cell) for cell in row))
    print()


def apply_clicks(N: int, c: int, X: list[list[int]]) -> list[list[int]]:
    """Board state produced by click matrix X from an all-zero start."""
    return [[(sum(X[r][j] for j in range(N))
              + sum(X[i][k] for i in range(N))
              - X[r][k]) % c
             for k in range(N)]
            for r in range(N)]


def verify_solution(N: int, c: int, target: list[list[int]],
                    X: list[list[int]]) -> bool:
    """Independently check X reaches `target`, without trusting the CNF."""
    reached = apply_clicks(N, c, X)
    for r in range(N):
        for k in range(N):
            if reached[r][k] != target[r][k]:
                print(f"  MISMATCH at ({r},{k}): "
                      f"reached={reached[r][k]}, target={target[r][k]}")
                return False
    return True


# ══════════════════════════════════════════════════════════════════════════
#  Timeout wrapper
# ══════════════════════════════════════════════════════════════════════════

_reporter = None
_quiet = 0


def report(progress: dict) -> None:
    """Send partial results to the parent, so a TIMEOUT row still has them."""
    if _reporter is not None and not _quiet:
        _reporter(progress)


@contextmanager
def quiet_reports():
    """Silence report() inside a nested solve (variant 2 called by variant 3)."""
    global _quiet
    _quiet += 1
    try:
        yield
    finally:
        _quiet -= 1


def _worker(solve_fn, N, c, target, conn):
    global _reporter
    _reporter = lambda progress: conn.send(("progress", progress))
    try:
        conn.send(("ok", solve_fn(N, c, target)))
    except Exception as exc:                     # noqa: BLE001
        conn.send(("error", repr(exc)))
    finally:
        conn.close()


def _stop(proc: Process) -> None:
    """Terminate a worker, escalating to SIGKILL if it ignores SIGTERM."""
    if proc.is_alive():
        proc.terminate()
        proc.join(5)
    if proc.is_alive():
        proc.kill()
        proc.join(5)


def run_with_timeout(solve_fn, N, c, target, timeout: float | None):
    """
    Run solve_fn(N, c, target) under a wall-clock limit.

    Returns (status, payload, elapsed, progress); status is "ok" | "timeout" |
    "error", and progress is the last partial state the solver sent via report().
    A child process, not a thread: the time is spent in Glucose's C extension,
    which no Python-level timer can interrupt. So solve_fn must be a
    module-level function and the caller needs an `if __name__` guard.

    A Pipe, not a Queue: Queue.put hands off to a feeder thread, which cannot run
    while Glucose holds the GIL, so progress sent just before a solve would never
    arrive. Pipe.send writes synchronously. The parent closes its copy of the
    write end, so a child that dies — even mid-message — surfaces as EOFError
    instead of a read that blocks forever.
    """
    if timeout is None:
        t0 = time.perf_counter()
        try:
            return "ok", solve_fn(N, c, target), time.perf_counter() - t0, {}
        except Exception as exc:                 # noqa: BLE001
            return "error", repr(exc), time.perf_counter() - t0, {}

    recv_end, send_end = Pipe(duplex=False)
    proc = Process(target=_worker, args=(solve_fn, N, c, target, send_end))
    t0 = time.perf_counter()
    deadline = t0 + timeout
    progress: dict = {}
    proc.start()
    send_end.close()

    try:
        while True:
            remaining = deadline - time.perf_counter()
            if remaining <= 0:
                _stop(proc)
                return "timeout", None, time.perf_counter() - t0, progress
            if not recv_end.poll(remaining):
                continue
            try:
                kind, payload = recv_end.recv()
            except (EOFError, OSError):
                _stop(proc)
                return ("error", "solver process exited without a result",
                        time.perf_counter() - t0, progress)

            if kind == "progress":
                progress.update(payload)
                continue
            elapsed = time.perf_counter() - t0
            _stop(proc)
            return kind, payload, elapsed, progress
    finally:
        recv_end.close()


# ══════════════════════════════════════════════════════════════════════════
#  Excel export
# ══════════════════════════════════════════════════════════════════════════

EXCEL_HEADERS = ["Instance", "N", "c", "Variables", "Clauses",
                 "Runtime (s)", "Timeout (s)", "SAT Calls", "Total Clicks",
                 "Optimal Clicks", "Status", "Machine"]


def export_to_excel(rows: list[dict], sheet_name: str, xlsx_path: str = "sat.xlsx",
                    quiet: bool = False, headers: list[str] | None = None):
    """
    Upsert rows into one sheet, keyed on the first header (Instance / Config).

    A key already in the sheet is overwritten in place, so re-running an
    instance replaces its old row instead of duplicating it. Headers missing
    from an existing sheet are added as new columns on the right.
    """
    headers = headers or EXCEL_HEADERS
    if os.path.exists(xlsx_path):
        wb = openpyxl.load_workbook(xlsx_path)
    else:
        wb = openpyxl.Workbook()
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(sheet_name)

    existing = [c.value for c in ws[1] if c.value is not None]
    for h in headers:
        if h not in existing:
            existing.append(h)
            ws.cell(row=1, column=len(existing), value=h).font = Font(bold=True)
    col_of = {h: i + 1 for i, h in enumerate(existing)}

    key_col = col_of[headers[0]]
    row_of = {ws.cell(row=r, column=key_col).value: r for r in range(2, ws.max_row + 1)}
    next_row = max(ws.max_row, 1) + 1

    for row in rows:
        r = row_of.get(row.get(headers[0]))
        if r is None:
            r, next_row = next_row, next_row + 1
        for h in headers:
            ws.cell(row=r, column=col_of[h], value=row.get(h, ""))

    os.makedirs(os.path.dirname(os.path.abspath(xlsx_path)), exist_ok=True)
    wb.save(xlsx_path)
    if not quiet:
        print(f"Results written to {xlsx_path} [{sheet_name}]")


def names_to_upgrade(xlsx_path: str, sheet_name: str, timeout: float) -> set[str]:
    """Keys of TIMEOUT rows whose recorded cutoff is below `timeout`."""
    if not os.path.exists(xlsx_path):
        return set()
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    if sheet_name not in wb.sheetnames:
        return set()
    rows = wb[sheet_name].iter_rows(values_only=True)
    header = list(next(rows))
    st, to = header.index("Status"), header.index("Timeout (s)")
    out = set()
    for r in rows:
        if not r[0] or str(r[st] or "").split(":")[0].strip() != "TIMEOUT":
            continue
        prev = r[to]
        if not isinstance(prev, (int, float)) or prev < timeout:
            out.add(str(r[0]))
    return out


def names_missing_fields(xlsx_path: str, sheet_name: str,
                         fields: list[str]) -> set[str]:
    """Keys of rows where any of `fields` is blank — for resuming a re-run."""
    if not os.path.exists(xlsx_path):
        return set()
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    if sheet_name not in wb.sheetnames:
        return set()
    rows = wb[sheet_name].iter_rows(values_only=True)
    header = list(next(rows))
    idx = [header.index(f) for f in fields if f in header]
    return {str(r[0]) for r in rows
            if r[0] and any(r[i] in (None, "") for i in idx)}


def names_with_status(xlsx_path: str, sheet_name: str, statuses: set[str]) -> set[str]:
    """Keys of rows whose Status is in `statuses` ("ERROR: …" matches ERROR)."""
    if not os.path.exists(xlsx_path):
        return set()
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    if sheet_name not in wb.sheetnames:
        return set()
    rows = wb[sheet_name].iter_rows(values_only=True)
    header = list(next(rows))
    st = header.index("Status")
    return {str(r[0]) for r in rows
            if r[0] and str(r[st] or "").split(":")[0].strip().upper() in statuses}


def already_done(xlsx_path: str, sheet_name: str) -> set[str]:
    """Instance names already recorded in this sheet, for --resume."""
    if not os.path.exists(xlsx_path):
        return set()
    try:
        wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    except (OSError, KeyError):
        return set()
    if sheet_name not in wb.sheetnames:
        return set()
    ws = wb[sheet_name]
    done = set()
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        if row[0]:
            done.add(str(row[0]))
    return done


# ══════════════════════════════════════════════════════════════════════════
#  Benchmark driver
# ══════════════════════════════════════════════════════════════════════════

def run_instance(solve_fn, inst: dict, timeout: float | None,
                 verbose: bool = True) -> dict:
    """Solve one instance and return a row ready for export_to_excel."""
    N, c, target = inst["N"], inst["c"], inst["target"]

    if verbose:
        print("=" * 64)
        print(f"Instance: {inst['name']}   (N={N}, c={c})")
        print_board(target, "  Target T")

    status, payload, elapsed, progress = run_with_timeout(
        solve_fn, N, c, target, timeout)

    row = {"Instance": inst["name"], "N": N, "c": c,
           "Runtime (s)": round(elapsed, 4),
           "Timeout (s)": timeout if timeout else "none",
           "Machine": machine_label()}

    if status in ("timeout", "error"):
        # Whatever the solver last reported before it was stopped. Total Clicks
        # here is the best solution found so far — an upper bound, not an optimum.
        best_X = progress.pop("_best_X", None)
        row.update(progress)
        row.setdefault("Total Clicks", "N/A")
        row["Status"] = "TIMEOUT" if status == "timeout" else f"ERROR: {payload}"
        if best_X is not None and not verify_solution(N, c, target, best_X):
            row["Status"] += " (incumbent failed verification)"
        if verbose:
            print(f"  {row['Status']} after {elapsed:.1f}s — "
                  f"{row.get('SAT Calls', 0)} SAT calls, "
                  f"best clicks {row['Total Clicks']}\n")
        return row

    result: SolveResult = payload
    row["Variables"] = result.n_vars
    row["Clauses"] = result.n_clauses
    row.update(result.extra)

    if result.X is None:
        row["Status"] = "UNSAT"
        row["Total Clicks"] = "N/A"
        if verbose:
            print("  Result: UNSATISFIABLE — no click matrix reaches this target.\n")
        return row

    total = sum(sum(r) for r in result.X)
    row["Total Clicks"] = total
    # Only a minimising run has proved an optimum; feasibility leaves this blank
    # rather than passing off "the solution we happened to find" as minimal.
    if result.optimum is not None:
        row["Optimal Clicks"] = result.optimum

    ok = verify_solution(N, c, target, result.X)
    row["Status"] = "OK" if ok else "VERIFY FAILED"

    if verbose:
        print_board(result.X, "  Solution X (click matrix)")
        print(f"  Encoding: {result.n_vars} vars, {result.n_clauses} clauses")
        print(f"  Total clicks: {total}")
        if result.extra:
            print(f"  Extra: {result.extra}")
        print(f"  Runtime: {elapsed:.4f}s")
        print(f"  Verification: {'PASSED' if ok else 'FAILED'}\n")

    return row


# ══════════════════════════════════════════════════════════════════════════
#  CLI
# ══════════════════════════════════════════════════════════════════════════

def build_parser(description: str) -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description=description)
    p.add_argument("--input", type=str, help="Path to one instance JSON")
    p.add_argument("--input-dir", type=str, help="Directory of instance JSONs")
    p.add_argument("--N", type=int, help="Grid size (N×N)")
    p.add_argument("--c", type=int, help="Number of colours / modulus")
    p.add_argument("--target", type=str, help='Target, e.g. "1,0,2;0,1,0;2,0,1"')
    p.add_argument("--xlsx", type=str, default="sat.xlsx", help="Excel output path")
    p.add_argument("--no-excel", action="store_true", help="Skip Excel export")
    p.add_argument("--timeout", type=float, default=600.0,
                   help="Per-instance wall-clock limit in seconds (0 = none)")
    p.add_argument("--quiet", action="store_true", help="Suppress per-instance output")
    p.add_argument("--upgrade-timeouts", action="store_true",
                   help="Re-run only TIMEOUT rows whose recorded cutoff is "
                        "below the current --timeout")
    p.add_argument("--only-missing", action="store_true",
                   help="Re-run only rows missing the logged solver fields")
    p.add_argument("--only-status", type=str,
                   help="Re-run only rows whose Status is one of these, "
                        "e.g. TIMEOUT,UNSAT")
    p.add_argument("--resume", action="store_true",
                   help="Skip instances already present in the target sheet")
    return p


def collect_instances(args) -> list[dict]:
    """Turn parsed CLI args into a list of instance dicts."""
    if args.input:
        return [load_instance(args.input)]
    if args.input_dir:
        return [load_instance(f) for f in find_instance_files(args.input_dir)]
    if args.N is not None and args.c is not None and args.target is not None:
        return [{"name": f"cli_{args.N}x{args.N}_c{args.c}",
                 "N": args.N, "c": args.c,
                 "target": parse_target(args.target, args.N, args.c)}]
    return []


def main(solve_fn, sheet_name: str, description: str = "Alien Tiles SAT solver"):
    """Generic CLI entry point, for a variant that only needs the defaults."""
    parser = build_parser(description)
    args = parser.parse_args()

    instances = collect_instances(args)
    if not instances:
        parser.print_help()
        sys.exit(1)

    timeout = args.timeout if args.timeout and args.timeout > 0 else None
    rows = [run_instance(solve_fn, inst, timeout, verbose=not args.quiet)
            for inst in instances]

    if not args.no_excel:
        export_to_excel(rows, sheet_name, args.xlsx)

    solved = sum(1 for r in rows if r.get("Status") == "OK")
    print(f"Done: {solved}/{len(rows)} solved and verified.")
